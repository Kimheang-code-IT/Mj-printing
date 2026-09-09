"""Synchronous PDF/XLSX renderers over report-service payloads."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.core.config import settings


def build_pdf(*, title: str, company: dict, period: dict, columns: list[str], rows: Iterable[dict], totals: dict | None = None) -> bytes:
    rows = list(rows)
    per_page = 38
    pages = [rows[i:i + per_page] for i in range(0, len(rows), per_page)] or [[]]
    buffer = BytesIO()
    document = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    for page_index, page_rows in enumerate(pages, 1):
        y = height - 42
        logo = str(company.get("logo") or "").replace("\\", "/").strip("/")
        if logo and ".." not in logo.split("/"):
            logo_path = Path(settings.local_storage_dir).resolve() / logo
            if logo_path.is_file():
                document.drawImage(str(logo_path), 42, y - 34, width=34, height=34, preserveAspectRatio=True, mask="auto")
        document.setFont("Helvetica-Bold", 14)
        document.drawString(84 if logo else 42, y, str(company.get("name") or ""))
        document.setFont("Helvetica", 8)
        document.drawString(42, y - 15, str(company.get("address") or ""))
        document.drawString(42, y - 27, str(company.get("phone") or ""))
        document.setFont("Helvetica-Bold", 12)
        document.drawString(42, y - 50, title)
        document.setFont("Helvetica", 8)
        document.drawString(42, y - 64, f"Period: {period.get('from', '')} to {period.get('to', '')}")
        document.setFont("Helvetica-Bold", 6)
        document.drawString(42, y - 84, " | ".join(columns)[:170])
        y -= 98
        document.setFont("Helvetica", 6)
        for row in page_rows:
            document.drawString(42, y, " | ".join(str(row.get(column, "")) for column in columns)[:180])
            y -= 15
        if page_index == len(pages) and totals:
            document.setFont("Helvetica-Bold", 7)
            document.drawString(42, max(45, y - 4), "Totals: " + ", ".join(f"{k}={v}" for k, v in totals.items())[:160])
        document.setFont("Helvetica", 8)
        document.drawCentredString(width / 2, 24, f"Page {page_index} of {len(pages)}")
        document.showPage()
    document.save()
    return buffer.getvalue()


def build_xlsx(*, title: str, company: dict, period: dict, columns: list[str], rows: Iterable[dict], totals: dict | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append([company.get("name", "")])
    sheet.append([title])
    sheet.append([f"Period: {period.get('from', '')} to {period.get('to', '')}"])
    sheet.append([])
    sheet.append(columns)
    for cell in sheet[5]:
        cell.font = Font(bold=True)
    money_columns = {name for name in columns if any(token in name for token in ("amount", "price", "cost", "profit", "debit", "credit", "balance", "value"))}
    for row in rows:
        values = []
        for column in columns:
            value = row.get(column, "")
            if isinstance(value, datetime) and value.tzinfo is not None:
                value = value.replace(tzinfo=None)
            values.append(value)
        sheet.append(values)
        row_number = sheet.max_row
        for index, column in enumerate(columns, 1):
            cell = sheet.cell(row=row_number, column=index)
            if column in money_columns and isinstance(cell.value, (Decimal, int, float)):
                cell.number_format = '#,##0.00'
            elif isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
    if totals:
        sheet.append([])
        sheet.append(["Totals"] + [f"{key}: {value}" for key, value in totals.items()])
    sheet.freeze_panes = "A6"
    for index, column in enumerate(columns, 1):
        width = max(len(column), *(len(str(sheet.cell(row=r, column=index).value or "")) for r in range(6, sheet.max_row + 1)))
        sheet.column_dimensions[get_column_letter(index)].width = min(40, width + 2)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
